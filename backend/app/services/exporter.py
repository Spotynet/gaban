"""Generación de archivos Excel (openpyxl) y PDF (reportlab) para reportes."""
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

_HEADER_FILL = PatternFill("solid", fgColor="4F46E5")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


def _sheet_from_rows(ws, headers, rows):
    ws.append(headers)
    for c in ws[1]:
        c.fill = _HEADER_FILL
        c.font = _HEADER_FONT
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, 1):
        width = max(len(str(h)), *(len(str(r[i - 1])) for r in rows)) if rows else len(str(h))
        ws.column_dimensions[chr(64 + i)].width = min(width + 3, 45)


def ventas_xlsx(ventas: list[dict], top: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ventas"
    _sheet_from_rows(ws,
        ["ID", "Serie-Folio", "Fecha", "Moneda", "Subtotal", "Descuento", "Impuesto", "Total", "Estado"],
        [[v["id"], f'{v.get("fiscal_serie") or ""}-{v.get("fiscal_folio") or ""}',
          v["created_at"][:16], v["currency"], v["subtotal"], v["discount"], v["tax"], v["total"], v["status"]]
         for v in ventas])
    ws2 = wb.create_sheet("Top productos")
    _sheet_from_rows(ws2, ["Producto", "SKU", "Unidades", "Ingresos"],
                     [[t["producto"], t["variant_sku"], t["unidades"], t["ingresos"]] for t in top])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def inventario_xlsx(rows: list[dict], resumen: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario"
    _sheet_from_rows(ws,
        ["SKU", "EAN", "Talla", "Color", "Cantidad", "Mínimo", "Costo unit.", "Valor costo"],
        [[r["variant_sku"], r.get("ean13") or "", r.get("size") or "", r.get("color") or "",
          r["quantity"], r["min_stock"], r["cost_price"], r["valor"]] for r in rows])
    ws2 = wb.create_sheet("Resumen")
    _sheet_from_rows(ws2, ["Métrica", "Valor"], [
        ["Valor a costo", resumen["valor_costo"]],
        ["Valor a precio venta", resumen["valor_venta"]],
        ["Líneas de stock", resumen["lineas"]],
        ["Bajo mínimo", resumen["bajo_minimo"]],
    ])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def resumen_pdf(titulo: str, periodo: str, kpis: dict, por_dia: list[dict]) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title=titulo)
    styles = getSampleStyleSheet()
    el = [Paragraph(titulo, styles["Title"]),
          Paragraph(periodo, styles["Normal"]), Spacer(1, 0.5 * cm)]

    kpi_rows = [["Indicador", "Valor"]] + [[k, str(v)] for k, v in kpis.items()]
    kt = Table(kpi_rows, colWidths=[9 * cm, 6 * cm])
    kt.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4F46E5")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F4F6")]),
    ]))
    el += [kt, Spacer(1, 0.6 * cm), Paragraph("Ventas por día", styles["Heading2"])]

    vd_rows = [["Día", "# Ventas", "Total"]] + [[d["dia"], str(d["ventas"]), f'{d["total"]:.2f}'] for d in por_dia]
    vt = Table(vd_rows, colWidths=[6 * cm, 4 * cm, 5 * cm])
    vt.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#059669")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
    ]))
    el.append(vt)
    doc.build(el)
    return buf.getvalue()
